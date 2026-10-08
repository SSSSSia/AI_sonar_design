# -*- coding: utf-8 -*-
"""军标格式 Word 导出（T11.5，P2-11）：IR2 设计文档 + 参数表 + 溯源依据 + 校验结论。

结构按总体设计 §5.2 排版约定：封面 → 设计文档正文（IR2.design_document，Markdown
轻量渲染：标题/表格/段落）→ 参数表（来源分类列）→ 溯源依据（source_type 四分类分组）
→ 校验结论（overall/CSR/issue 分级）。

如实标注（A-16）：内置程序化模板（黑体标题/宋体正文/表格样式），**非授权军标正式
模板**；`DOCX_TEMPLATE_PATH` 指向 .docx 模板文件时以其样式为底覆挂（占位替换留待
正式模板到位后按需扩展）。
"""

from __future__ import annotations

import io
import os
import re
from typing import Any

SOURCE_TYPE_LABEL = {
    "formula": "公式计算",
    "case": "案例",
    "standard": "军标",
    "estimated": "约束内提议",
}
SEVERITY_LABEL = {"error": "错误", "warning": "警告", "info": "提示", "unverified": "未验证"}


def _docx():
    import docx  # python-docx（requirements 既有，T3.1 文档解析同款）

    return docx


def _set_base_styles(doc: Any) -> None:
    from docx.shared import Pt

    for name, size in (("Normal", 10.5), ("Title", 22), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        try:
            st = doc.styles[name]
            st.font.size = Pt(size)
            st.font.name = "宋体" if name == "Normal" else "黑体"
        except KeyError:
            continue


def _add_cover(doc: Any, title: str, meta: list[tuple[str, str]]) -> None:
    doc.add_heading(title, level=0)
    table = doc.add_table(rows=len(meta), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(meta):
        table.cell(i, 0).text = k
        table.cell(i, 1).text = v
    doc.add_page_break()


_MD_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")


def _add_markdown_body(doc: Any, markdown: str) -> None:
    """轻量 Markdown 渲染：#/##/### → 标题；|表格| → docx 表；其余 → 段落。

    不追求完整 MD 语义（代码块/嵌套列表按段落如实呈现）——正文素材由 A2 生成，
    结构相对规整（A-16：渲染保真度如实记录）。
    """
    lines = (markdown or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        m = _MD_TABLE_ROW.match(line)
        if m and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").replace(" ", "")) <= {"-", ":"} \
                and lines[i + 1].strip().startswith("|"):
            header = [c.strip() for c in m.group(1).split("|")]
            rows: list[list[str]] = [header]
            j = i + 2
            while j < len(lines):
                m2 = _MD_TABLE_ROW.match(lines[j])
                if not m2:
                    break
                rows.append([c.strip() for c in m2.group(1).split("|")])
                j += 1
            table = doc.add_table(rows=len(rows), cols=len(header))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row[: len(header)]):
                    table.cell(r, c).text = cell
            i = j
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        else:
            doc.add_paragraph(stripped)
        i += 1


def _add_param_table(doc: Any, params: list[dict[str, Any]]) -> None:
    doc.add_heading("参数表", level=1)
    if not params:
        doc.add_paragraph("（无登记参数）")
        return
    table = doc.add_table(rows=len(params) + 1, cols=7)
    table.style = "Light Grid Accent 1"
    headers = ["参数ID", "名称", "数值", "单位", "来源分类", "来源依据", "置信度"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h
    for r, p in enumerate(params, 1):
        values = [str(p.get("id") or ""), str(p.get("name") or ""), str(p.get("value") or ""),
                  str(p.get("unit") or ""), SOURCE_TYPE_LABEL.get(str(p.get("source_type")), str(p.get("source_type") or "")),
                  str(p.get("source_ref") or ""), str(p.get("confidence") or "")]
        for c, v in enumerate(values):
            table.cell(r, c).text = v


def _add_trace_chapter(doc: Any, params: list[dict[str, Any]]) -> None:
    doc.add_heading("溯源依据", level=1)
    groups: dict[str, list[dict[str, Any]]] = {}
    for p in params or []:
        groups.setdefault(str(p.get("source_type") or "estimated"), []).append(p)
    if not groups:
        doc.add_paragraph("（无登记参数，无溯源项）")
        return
    for stype in ("formula", "standard", "case", "estimated"):
        items = groups.get(stype)
        if not items:
            continue
        doc.add_heading(f"{SOURCE_TYPE_LABEL.get(stype, stype)}（{len(items)} 项）", level=2)
        for p in items:
            doc.add_paragraph(
                f"{p.get('name')}（{p.get('id')}）= {p.get('value')} {p.get('unit')} —— "
                f"依据：{p.get('source_ref') or '—'}；置信度：{p.get('confidence')}",
                style="List Bullet",
            )
            derivation = str(p.get("derivation") or "").strip()
            if derivation:
                doc.add_paragraph(f"计算/引用说明：{derivation}")


def _add_validation_chapter(doc: Any, ir4: dict[str, Any] | None) -> None:
    doc.add_heading("校验结论", level=1)
    if not ir4:
        doc.add_paragraph("IR4 校验报告尚未生成（如实说明）。")
        return
    overall = str(ir4.get("overall") or "unknown")
    overall_text = {"pass": "通过", "fail": "未通过",
                    "incomplete": "不完整（存在应检未检项）"}.get(overall, overall)
    doc.add_paragraph(f"总体结论：{overall_text}")
    csr = ir4.get("csr")
    if csr is not None:
        doc.add_paragraph(f"约束满足率（CSR）：{csr} / 100")
    coverage = ir4.get("coverage")
    if coverage:
        ratio = coverage.get("ratio")
        ratio_text = "—" if ratio is None else f"{ratio}%"
        doc.add_paragraph(f"验证覆盖率：{ratio_text}"
                          f"（已检 {coverage.get('checked', 0)} / 适用必检 {coverage.get('applicable', 0)} 项；"
                          "必检项=公式复算+适用军标条目+几何联动期望，未验证项计入分母）")
    issues = ir4.get("issues") or []
    if not issues:
        doc.add_paragraph("无校验问题项。")
        return
    table = doc.add_table(rows=len(issues) + 1, cols=4)
    table.style = "Light Grid Accent 1"
    for c, h in enumerate(["约束ID", "级别", "说明", "建议"]):
        table.cell(0, c).text = h
    for r, it in enumerate(issues, 1):
        sev = str(it.get("severity") or "")
        values = [str(it.get("constraint_id") or ""), SEVERITY_LABEL.get(sev, sev),
                  str(it.get("message") or "—"), str(it.get("suggestion") or "—")]
        for c, v in enumerate(values):
            table.cell(r, c).text = v


def export_session_docx(session: dict[str, Any], ir2: dict[str, Any] | None,
                        ir3: dict[str, Any] | None, ir4: dict[str, Any] | None,
                        ir2_version: int) -> bytes:
    """会话交付文档 → docx 字节流。素材缺失各章如实标注，不造假。"""
    docx = _docx()
    template_path = os.environ.get("DOCX_TEMPLATE_PATH", "")
    if template_path and os.path.exists(template_path):
        doc = docx.Document(template_path)
    else:
        doc = docx.Document()
    _set_base_styles(doc)

    title = str(session.get("title") or "声呐装备设计方案")
    _add_cover(doc, title, [
        ("项目", str(session.get("project_name") or "—")),
        ("会话ID", str(session.get("id") or "")),
        ("IR2 版本", f"v{ir2_version}"),
        ("生成方式", "平台自动排版（内置模板；正式军标模板待授权后替换）"),
    ])
    doc.add_heading("设计文档", level=1)
    if ir2:
        _add_markdown_body(doc, str(ir2.get("design_document") or ""))
    else:
        doc.add_paragraph("IR2 尚未生成。")
    params = list((ir2 or {}).get("parameters") or [])
    _add_param_table(doc, params)
    _add_trace_chapter(doc, params)
    _add_validation_chapter(doc, ir4)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
