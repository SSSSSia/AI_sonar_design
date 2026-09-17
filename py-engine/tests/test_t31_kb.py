# -*- coding: utf-8 -*-
"""T3.1 知识文档上传与解析单测：五格式解析器（含 OCR 降级）+ 上传 API 行为 + 鉴权。"""

import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from kb import parsers


# ----------------------------------------------------------------------
# 解析器
# ----------------------------------------------------------------------

def test_parse_txt_paragraphs_with_location():
    data = "第一段内容。\n\n第二段内容，较长一些。\n\n\n第三段。".encode("utf-8")
    result = parsers.parse_document("样例.txt", data)
    assert [u.location for u in result.units] == ["段落1", "段落2", "段落3"]
    assert result.units[0].text == "第一段内容。"
    assert result.ocr_pending == 0


def test_parse_md_chapter_paths():
    data = (
        "文件级引言。\n"
        "# 拖曳阵知识\n\n前言内容。\n"
        "## 声学约束\n\nDI≈10lgN。\n"
        "### 间距\n\n间距 1.5m。\n"
        "## 尺寸约束\n\n外径≤27mm。\n"
    ).encode("utf-8")
    result = parsers.parse_document("kb.md", data)
    locations = [u.location for u in result.units]
    assert "（前言）" in locations                    # 首个标题之前的文本
    assert "拖曳阵知识" in locations                 # # 级章节
    assert "拖曳阵知识 > 声学约束" in locations
    assert "拖曳阵知识 > 声学约束 > 间距" in locations
    assert "拖曳阵知识 > 尺寸约束" in locations
    assert result.section_count == 4


def test_parse_docx_headings_and_paragraphs():
    import io

    import docx

    doc = docx.Document()
    doc.add_heading("系统组成", level=1)
    doc.add_paragraph("拖曳线列阵由声阵段等组成。")
    doc.add_heading("约束", level=2)
    doc.add_paragraph("外径不超过 27mm。")
    buf = io.BytesIO()
    doc.save(buf)
    result = parsers.parse_document("手册.docx", buf.getvalue())
    assert len(result.units) == 2
    assert result.units[0].location.startswith("系统组成#段落")
    assert result.units[1].location.startswith("系统组成 > 约束#段落")
    assert "外径不超过 27mm" in result.units[1].text


def test_parse_xlsx_rows_with_sheet_location():
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "参数"
    ws.append(["参数", "取值"])
    ws.append(["阵元总数", 16])
    ws.append([None, None])  # 空行跳过
    ws.append(["阵元间距", 1.5])
    buf = io.BytesIO()
    wb.save(buf)
    result = parsers.parse_document("参数表.xlsx", buf.getvalue())
    locations = [u.location for u in result.units]
    assert locations == ["参数!行1", "参数!行2", "参数!行4"]
    assert "16" in result.units[1].text


def _make_pdf(text_lines: list[str]) -> bytes:
    """构造最小合法 PDF（Helvetica，仅 ASCII 文本；供解析与 OCR 降级测试）。"""
    content = "BT /F1 18 Tf 72 720 Td " + " ".join(
        f"({line}) Tj 0 -24 Td" for line in text_lines
    ) + " ET"
    content_obj = f"<< /Length {len(content)} >>\nstream\n{content}\nendstream"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        content_obj,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    return bytes(out)


def test_parse_pdf_text_page_with_location():
    result = parsers.parse_document("手册.pdf", _make_pdf(["towed array knowledge", "DI = 10 lg N"]))
    assert result.page_count == 1
    assert result.ocr_pending == 0
    assert result.units[0].location == "第1页"
    assert "array" in result.units[0].text


def test_parse_pdf_blank_pages_counted_as_ocr_pending():
    """扫描页（空文本页）如实计入 ocr_pending——OCR 插件待接（D-14），不伪造文本。"""
    import io

    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    result = parsers.parse_document("扫描件.pdf", buf.getvalue())
    assert result.page_count == 2
    assert result.ocr_pending == 2
    assert result.units == []


def test_unsupported_format_rejected():
    with pytest.raises(parsers.UnsupportedFormatError):
        parsers.parse_document("病毒.exe", b"MZ...")


# ----------------------------------------------------------------------
# 上传 API（fake 存储层）
# ----------------------------------------------------------------------
@pytest.fixture()
def client(monkeypatch):
    from api.main import app

    return TestClient(app)


@pytest.fixture()
def fake_store(monkeypatch):
    """替换 kb_store 落库与 MinIO 上传（API 行为测试不触真实存储）。"""
    rows: dict = {}

    def fake_create(filename, fmt, size_bytes, bucket, object_key, uploaded_by):
        doc = {"doc_id": "abc1234567890def", "filename": filename, "format": fmt,
               "size_bytes": size_bytes, "bucket": bucket, "object_key": object_key,
               "status": "uploaded", "ocr_pending": 0}
        rows["abc1234567890def"] = doc
        return doc

    monkeypatch.setattr("storage.kb_store.create_document", fake_create)
    monkeypatch.setattr("storage.kb_store.update_document",
                        lambda doc_id, **fields: rows[doc_id].update(fields))
    monkeypatch.setattr("storage.kb_store.get_document", lambda doc_id: rows.get(doc_id))
    monkeypatch.setattr("storage.kb_store.list_documents", lambda limit=100: list(rows.values()))
    monkeypatch.setattr("storage.kb_store.delete_document", lambda doc_id: rows.pop(doc_id, None) is not None)
    monkeypatch.setattr("storage.object_store.upload_bytes",
                        lambda data, bucket, key, content_type="": {"bucket": bucket, "object_key": key})
    monkeypatch.setattr("storage.object_store.remove_object", lambda bucket, key: None)
    return rows


AUTH = {"X-User-Id": "tester", "X-User-Roles": "admin"}


def test_upload_txt_happy_path(client, fake_store):
    r = client.post("/kb/documents", headers=AUTH,
                    files={"file": ("样例.txt", "第一段。\n\n第二段。".encode("utf-8"), "text/plain")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "parsed" and body["units"] == 2
    assert fake_store["abc1234567890def"]["status"] == "parsed"


def test_upload_requires_writer_role(client, fake_store):
    r = client.post("/kb/documents", headers={"X-User-Id": "viewer", "X-User-Roles": "viewer"},
                    files={"file": ("a.txt", b"x", "text/plain")})
    assert r.status_code == 403
    r = client.post("/kb/documents", files={"file": ("a.txt", b"x", "text/plain")})
    assert r.status_code == 401


def test_upload_bad_format_and_empty(client, fake_store):
    r = client.post("/kb/documents", headers=AUTH,
                    files={"file": ("a.exe", b"MZ", "application/octet-stream")})
    assert r.status_code == 400
    r = client.post("/kb/documents", headers=AUTH,
                    files={"file": ("empty.txt", b"", "text/plain")})
    assert r.status_code == 400


def test_upload_oversized_413(client, fake_store, monkeypatch):
    import app_config

    monkeypatch.setattr(app_config, "KB_UPLOAD_MAX_BYTES", 10)
    r = client.post("/kb/documents", headers=AUTH,
                    files={"file": ("big.txt", b"x" * 100, "text/plain")})
    assert r.status_code == 413


def test_parse_failure_marks_failed_not_raise(client, fake_store, monkeypatch):
    """解析失败不中断服务（§6 ①）：doc=failed + error 如实，接口仍 200。"""
    monkeypatch.setattr(parsers, "parse_document",
                        lambda filename, data: (_ for _ in ()).throw(RuntimeError("坏文档")))
    r = client.post("/kb/documents", headers=AUTH,
                    files={"file": ("坏.txt", b"data", "text/plain")})
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "failed" and "坏文档" in body["error"]


def test_delete_document_cleans_up(client, fake_store):
    client.post("/kb/documents", headers=AUTH,
                files={"file": ("样例.txt", "内容".encode("utf-8"), "text/plain")})
    r = client.delete("/kb/documents/abc1234567890def", headers=AUTH)
    assert r.status_code == 200 and r.json()["deleted"] is True
    r = client.delete("/kb/documents/abc1234567890def", headers=AUTH)
    assert r.status_code == 404
